"""Étape 5 - Données livrées : GeoPackage (QGIS), géodatabase fichier (ArcGIS Pro), rasters,
classeur Excel des tableaux de l'état initial.

    python 05_donnees_livrees.py
"""
import shutil

import config
import geopandas as gpd
import pandas as pd
import pyogrio
import shapely
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from etat_initial.reseau import wfs
from etat_initial.sources import meme_dimension
from outils import ecrire, log

SIG = config.LIVRABLES / "02_Donnees_SIG"
GPKG = SIG / "etat_initial.gpkg"
RASTERS = SIG / "rasters"
TABLEUR = config.LIVRABLES / "04_Tableaux" / "Etat_initial_tableaux.xlsx"

# couches du GeoPackage de travail reprises telles quelles (nom livré = nom de travail)
COUCHES = ["projet", "aire_etude", "commune", "znieff1", "znieff2", "zsc", "zps", "rnr", "rnn", "cen",
           "ramsar", "pnr", "zdh", "cours_eau", "plan_eau", "clc", "rpg", "foret", "haie",
           "unesco_bien", "unesco_tampon", "terril_unesco", "pprt", "ppr_inond", "ppr_minier", "nappe",
           "cavite", "sis", "ssp", "ligne_elec", "poste_elec", "parc_pv", "eolienne", "route",
           "voie_ferree", "batiment", "doc_urba", "zone_urba", "prescription_surf", "prescription_lin",
           "prescription_pct", "info_surf", "assiette_sup_s", "assiette_sup_l", "assiette_sup_p",
           "generateur_sup_s", "generateur_sup_l", "generateur_sup_p", "parcelle", "parcelle_zip",
           "ocsge", "ocsge_evolution", "observation_patrimoniale", "enjeu"]
TABLES = {"t_enjeux": "Enjeux", "t_zonages": "Zonages", "t_zones_humides": "Zones humides",
          "t_occupation": "Occupation du sol", "t_ocs_sources": "Trois sources OCS",
          "t_ocs_accords": "Concordance OCS", "t_zonage": "Zonage PLU", "t_prescriptions": "Prescriptions PLU",
          "t_servitudes": "Servitudes", "t_foncier": "Foncier", "t_risques": "Risques",
          "t_reseaux": "Réseaux et cumuls", "t_especes_groupes": "Espèces par groupe",
          "t_especes": "Espèces"}


def geopackage():
    log(f"GeoPackage livré : {GPKG.relative_to(config.RACINE)}")
    if GPKG.exists():
        GPKG.unlink()
    presentes = {n for n, _ in pyogrio.list_layers(config.ETUDE)}
    for nom in COUCHES:
        if nom in presentes:
            gdf = gpd.read_file(config.ETUDE, layer=nom)
            if not gdf.empty:
                ecrire(gdf, nom, GPKG)
    for nom in TABLES:
        if nom in presentes:
            ecrire(pyogrio.read_dataframe(config.ETUDE, layer=nom, read_geometry=False), nom, GPKG)
    # contexte et criblage, depuis le référentiel territorial
    for nom, livre in (("commune", "contexte_commune"), ("epci", "contexte_epci"),
                       ("criblage_friches", "criblage_friches")):
        ecrire(gpd.read_file(config.REFERENTIEL, layer=nom), livre, GPKG)
    ecrire(pyogrio.read_dataframe(config.REFERENTIEL, layer="criblage_regles", read_geometry=False),
           "criblage_regles", GPKG)
    # carte de situation : départements de la région et point du site
    dep = wfs("geopf", "ADMINEXPRESS-COG.LATEST:departement", cql="code_insee_de_la_region = '32'")
    ecrire(dep[["code_insee", "nom_officiel", "geometry"]].rename(columns={"nom_officiel": "nom"}),
           "contexte_departement", GPKG)
    site = gpd.read_file(config.ETUDE, layer="projet")
    ecrire(site.assign(geometry=site.geometry.representative_point())[["nom", "geometry"]],
           "contexte_site", GPKG)


def rasters():
    RASTERS.mkdir(parents=True, exist_ok=True)
    for nom in ("oso_site.tif", "concordance_ocs.tif", "ocs_harmonise_ocsge.tif", "ocs_harmonise_clc.tif",
                "ocs_harmonise_oso.tif"):
        src = config.TRAVAIL / nom
        if src.exists():
            shutil.copy2(src, RASTERS / nom)
            log(f"  -> rasters/{nom}")


def geodatabase():
    """Géodatabase fichier pour les partenaires équipés d'ArcGIS (pilote OpenFileGDB de GDAL)."""
    gdb = SIG / "etat_initial.gdb"
    if gdb.exists():
        shutil.rmtree(gdb)
    log("Géodatabase fichier (ArcGIS Pro)")
    for nom, _ in pyogrio.list_layers(GPKG):
        gdf = gpd.read_file(GPKG, layer=nom) if nom not in TABLES and not nom.startswith("criblage_r") \
            else pyogrio.read_dataframe(GPKG, layer=nom, read_geometry=False)
        options = {}
        if isinstance(gdf, gpd.GeoDataFrame):
            # une classe d'entités n'accepte qu'un type de géométrie : tout est promu en multi
            # de la dimension dominante (les fragments d'une autre dimension sont écartés)
            dims = shapely.get_dimensions(gdf.geometry.values)
            dim = int(pd.Series(dims[dims >= 0]).mode().iloc[0])
            multi = [shapely.MultiPoint, shapely.MultiLineString, shapely.MultiPolygon][dim]
            geoms = []
            for g in gdf.geometry:
                g = meme_dimension(g, dim)
                geoms.append(None if g is None else g if g.geom_type.startswith("Multi") else multi([g]))
            gdf = gdf.set_geometry(gpd.GeoSeries(geoms, crs=gdf.crs))
            gdf = gdf[gdf.geometry.notna()]
            options["geometry_type"] = ["MultiPoint", "MultiLineString", "MultiPolygon"][dim]
        pyogrio.write_dataframe(gdf, gdb, layer=nom, driver="OpenFileGDB", **options)
    log(f"  -> {gdb.name} : {len(pyogrio.list_layers(gdb))} classes et tables")


ENTETES = {
    "theme": "Thème", "sujet": "Sujet", "niveau": "Niveau", "justification": "Justification",
    "regle": "Règle appliquée", "code_couche": "Couche", "identifiant": "Identifiant",
    "distance_zip_m": "Distance à la ZIP (m)", "part_zip": "Part de la ZIP couverte", "surface_ha": "Surface (ha)",
    "surface_totale_ha": "Surface totale (ha)", "url": "Fiche", "code_cs": "Code couverture",
    "code_us": "Code usage", "part": "Part de l'aire", "nb_objets": "Nombre d'objets",
    "nb_parcelles": "Parcelles", "surface_zip_ha": "Surface dans la ZIP (ha)", "proprietaire": "Propriétaire",
    "nom_scientifique": "Nom scientifique", "nom_francais": "Nom français", "nb_donnees": "Données",
    "protegee": "Protégée", "liste_rouge_nationale": "Liste rouge France",
    "liste_rouge_regionale": "Liste rouge régionale", "determinante_znieff": "Déterminante ZNIEFF",
    "patrimoniale": "Patrimoniale", "especes": "Espèces", "protegees": "Protégées",
    "patrimoniales": "Patrimoniales", "donnees": "Données", "accord": "Accord", "kappa": "Kappa",
    "comparaison": "Comparaison", "intitule": "Intitulé", "categorie": "Catégorie", "libelle": "Libellé",
    "present": "Présence", "detail": "Détail", "objet": "Objet", "nombre": "Nombre", "milieu": "Milieu",
}


def tableur():
    log(f"Classeur : {TABLEUR.relative_to(config.RACINE)}")
    TABLEUR.parent.mkdir(parents=True, exist_ok=True)
    entete = PatternFill("solid", fgColor="1D3557")
    with pd.ExcelWriter(TABLEUR, engine="openpyxl") as xl:
        for nom, onglet in TABLES.items():
            df = pyogrio.read_dataframe(GPKG, layer=nom, read_geometry=False)
            df = df.drop(columns=[c for c in ("rang_niveau", "species_key", "cd_nom") if c in df])
            df = df.rename(columns=lambda c: ENTETES.get(c, c[:1].upper() + c[1:].replace("_", " ")))
            df.to_excel(xl, sheet_name=onglet[:31], index=False)
            ws = xl.sheets[onglet[:31]]
            ws.freeze_panes = "A2"
            for i, col in enumerate(df.columns, start=1):
                c = ws.cell(row=1, column=i)
                c.font, c.fill = Font(bold=True, color="FFFFFF"), entete
                c.alignment = Alignment(wrap_text=True, vertical="center")
                largeur = min(70, max(10, int(df[col].astype(str).str.len().quantile(0.9) if len(df) else 10)))
                ws.column_dimensions[get_column_letter(i)].width = largeur
            ws.auto_filter.ref = ws.dimensions


if __name__ == "__main__":
    geopackage()
    rasters()
    geodatabase()
    tableur()
